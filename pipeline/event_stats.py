"""事件統計圖:一段期間(例如颱風影響期間)各測站的總累積雨量、最大平均風、最大陣風 → PNG + CSV。
色階、風級與網頁相同(雨量用網頁 ECMWF 長時累積的色階,風用蒲福風級);底圖用 shp/TWN_CITY(--shp 指定資料夾)。
  python -m pipeline.event_stats --start "2026-10-22 00:00" [--end "2026-10-25 12:00"] [--local]
時間為台灣時間;--end 省略 = 最新一筆快照。輸出到 charts/{起}_{迄}/。
本機 data/stations 沒有的快照自動從 R2 下載(存在 pipeline/.cache/stations/),任何一台電腦都能跑;
--local 只用本機資料。

總累積雨量:每站把期間切成首尾相接的幾段,每段用某一筆快照的累積雨量欄位
(10 分鐘、1、3、6、12、24 小時滾動累積;今日、近 2 日、近 3 日的日曆日累積),段數最少的組合相加。
快照缺漏時會改用較長的累積時段跨過去;接不起來的站不畫(log 會列出數量)。
雨量欄位的時間 = 快照時間(與網頁、station_agg 相同)。
注意 CWA 的 Past2days / Past3days(rain_48h / rain_72h)是日曆日(昨天起 / 前天起,台灣 00 時),不是滾動 48 / 72 小時。
最大平均風 / 最大陣風:觀測時間落在期間內的各筆取最大;陣風時間用 gust_at(有的話)。
最大時雨量 / 3 / 6 / 12 / 24 小時雨量(rain_max.csv、縣市表):各快照的滾動累積取最大,只算整段都在期間內的。
颱風、豪雨事件一次產統計 + 影片 + 摘要:用 pipeline.event_report。"""
import argparse
from pathlib import Path
import csv
import json
from datetime import datetime, timedelta, timezone

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import requests
import shapely
from shapely.geometry import shape
from matplotlib.collections import PatchCollection
from matplotlib.cm import ScalarMappable
from matplotlib.colors import BoundaryNorm, ListedColormap
from matplotlib.patches import Polygon as MplPolygon

from . import config, idw
from .common import read_index

TW = timezone(timedelta(hours=8))
R2_URL = "https://pub-95cef959cf924e2e938a1d1b2a833e45.r2.dev/"
OUT_DIR = config.ROOT / "charts"
FONT = ["Microsoft JhengHei", "Noto Sans CJK TC", "PingFang TC"]

# 與網頁 web/js/colormaps.js、layers/ecmwf.js 相同
RAIN_LEVELS = [0,10, 20, 60, 100, 150, 200, 300, 400, 500, 600, 700, 800, 900, 1000, 1200, 1500, 2000]
RAIN_COLORS = ["#c1c2c2", "#9bffff", "#00cfff", "#0198ff", "#0165ff", "#309901", "#32ff00", "#f8ff00",
               "#ffcb00", "#ff9a00", "#fa0300", "#cc0003", "#a00000", "#98009a", "#c304cc", "#f805f3", "#fecbff"]
BEAUFORT = [0, 0.3, 1.6, 3.4, 5.5, 8.0, 10.8, 13.9, 17.2, 20.8, 24.5, 28.5, 32.7, 37.0, 41.5, 46.2, 51.0, 56.1, 61.3]
# 最後一個 61.3 = 17 級上界;61.3 以上另用 WIND_OVER,色階標 >17
WIND_COLORS = ["#eef5fb", "#c6dbef", "#9ecae1", "#6baed6", "#2171b5", "#74c476", "#238b45",
               "#fff59d", "#ffe03a", "#f5b700", "#ef5350", "#b71c1c", "#ce93d8", "#9c27b0",
               "#6a1b9a", "#4a148c", "#2a0a4a", "#14001f"]
WIND_OVER = "#7a7a7a"                           # 61.3 m/s 以上(超過 17 級)
LIGHT = {"bg": "#ffffff", "land": "#f5f5f5", "coast": "#000000", "county": "#1a1a1a"}

# 滾動累積的雨量欄位 → 分鐘數
RAIN_WIN = {"rain_10m": 10, "rain_1h": 60, "rain_3h": 180, "rain_6h": 360, "rain_12h": 720, "rain_24h": 1440}
# 日曆日累積(台灣時間 00 時起算)→ 往前幾天:今日 0、Past2days 昨天起、Past3days 前天起。
# 實測 Past2days 在 00:00 從 640 掉到 558(同時 24 小時 559→558),不是滾動 48 小時。
RAIN_DAYS = {"rain_today": 0, "rain_48h": 1, "rain_72h": 2}
WX = ["ws", "wd", "gust", "gust_dir"]

MAIN = (119.25, 122.15, 21.85, 25.7)           # 本島 + 澎湖、蘭嶼、綠島、彭佳嶼
INSETS = {"馬祖": (119.88, 120.55, 25.92, 26.42), "金門": (118.18, 118.52, 24.36, 24.55)}   # 小圖由上而下
REGION = (118.0, 122.3, 21.8, 26.5)            # 納入統計的範圍(排除日本氣象廳站等)


def stamp_dt(s):
    return datetime.strptime(s, "%Y%m%d%H%M").replace(tzinfo=timezone.utc)


def to_stamp(dt):
    return dt.astimezone(timezone.utc).strftime("%Y%m%d%H%M")


def tw(s, fmt="%Y/%m/%d %H:%M"):
    return stamp_dt(s).astimezone(TW).strftime(fmt)


def parse_tw(text):
    return datetime.strptime(text, "%Y-%m-%d %H:%M").replace(tzinfo=TW)


# ── 讀快照 ──
def snapshot_times(use_r2):
    local = set(read_index("stations").get("times", []))
    if use_r2:
        H = {"User-Agent": "Mozilla/5.0"}
        local |= set(requests.get(R2_URL + "stations/index.json", headers=H, timeout=60).json()["times"])
    return sorted(local)


def load_snapshot(stamp, use_r2):
    path = config.DATA_DIR / "stations" / f"{stamp}.json"
    if not path.exists():
        if not use_r2:
            return None
        path = config.CACHE_DIR / "stations" / f"{stamp}.json"
        if not path.exists():
            r = requests.get(f"{R2_URL}stations/{stamp}.json", headers={"User-Agent": "Mozilla/5.0"}, timeout=60)
            if r.status_code != 200:
                return None
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(r.content)
    return json.loads(path.read_text(encoding="utf-8"))


def collect(start, end, use_r2):
    """期間內(雨量多讀 start 之後;風讀觀測時間在 [start, end])的快照 → 站資料與數值陣列。"""
    times = [t for t in snapshot_times(use_r2) if start < t <= end]
    meta, col = {}, {}
    snaps = []
    for i, t in enumerate(times):
        snap = load_snapshot(t, use_r2)
        if snap is None:
            continue
        f = {k: j for j, k in enumerate(snap["fields"])}
        rows = []
        for r in snap["rows"]:
            if "type" in f and r[f["type"]] == "jma":
                continue
            lon, lat = r[f["lon"]], r[f["lat"]]
            if lon is None or not (REGION[0] <= lon <= REGION[1] and REGION[2] <= lat <= REGION[3]):
                continue
            sid = r[f["id"]]
            meta[sid] = {"id": sid, "name": r[f["name"]], "county": r[f.get("county")] if "county" in f else "",
                         "town": r[f["town"]] if "town" in f else "", "lon": lon, "lat": lat,
                         "alt": r[f["alt"]] if "alt" in f else None}
            col.setdefault(sid, len(col))
            rows.append(r)
        snaps.append((t, f, rows))
        if (i + 1) % 50 == 0:
            print(f"  讀取快照 {i + 1}/{len(times)}")
    return times, snaps, meta, col


# ── 總累積雨量 ──
def total_rain(start, end, snaps, col):
    """每站段數最少的首尾相接組合;回傳 (總雨量, 段數),接不起來為 NaN。"""
    n = len(col)
    done = {start: (np.zeros(n), np.zeros(n))}          # 時間 → (段數, 累積);段數 inf = 到不了
    last = start
    for t, f, rows in snaps:
        ids = np.array([col[r[f["id"]]] for r in rows], dtype=int)
        best_n, best_v = np.full(n, np.inf), np.zeros(n)
        tdt = stamp_dt(t)
        mid = tdt.astimezone(TW).replace(hour=0, minute=0)
        wins = [(k, to_stamp(tdt - timedelta(minutes=m))) for k, m in RAIN_WIN.items() if k in f]
        if mid != tdt.astimezone(TW):          # 00:00 那筆的日累積還是前一天的(今日 = 前一天整天),不用
            wins += [(k, to_stamp(mid - timedelta(days=d))) for k, d in RAIN_DAYS.items() if k in f]
        for k, origin in wins:
            if origin < start or origin not in done:
                continue
            on, ov = done[origin]
            x = np.full(n, np.nan)
            x[ids] = [np.nan if r[f[k]] is None else r[f[k]] for r in rows]
            cn = np.where(np.isfinite(x), on + 1, np.inf)
            better = cn < best_n
            best_n = np.where(better, cn, best_n)
            best_v = np.where(better, ov + np.nan_to_num(x), best_v)
        done[t] = (best_n, best_v)
        last = t
    seg, val = done[last]
    return np.where(np.isfinite(seg), val, np.nan), seg, last


# ── 最大平均風 / 最大陣風 ──
def max_wind(start, end, snaps, col):
    n = len(col)
    out = {v: {"val": np.full(n, np.nan), "dir": np.full(n, np.nan), "at": [None] * n} for v in ("ws", "gust")}
    for t, f, rows in snaps:
        for r in rows:
            obs = r[f["obs"]] or ""
            if not (start <= obs <= end):
                continue
            i = col[r[f["id"]]]
            for v, d in (("ws", "wd"), ("gust", "gust_dir")):
                x = r[f[v]] if v in f else None
                if x is None:
                    continue
                at = r[f["gust_at"]] if v == "gust" and "gust_at" in f and r[f["gust_at"]] else obs
                if v == "gust" and not (start <= at <= end):
                    continue
                o = out[v]
                cur = o["val"][i]
                if np.isnan(cur) or x > cur or (x == cur and at < (o["at"][i] or "")):
                    o["val"][i], o["at"][i] = x, at
                    o["dir"][i] = r[f[d]] if d in f and r[f[d]] is not None else np.nan
    return out


# ── 最大短時雨量(豪雨事件用)──
RAIN_MAX = {"rain_1h": "最大時雨量", "rain_3h": "最大3小時雨量", "rain_6h": "最大6小時雨量",
            "rain_12h": "最大12小時雨量", "rain_24h": "最大24小時雨量"}


def max_rain(start, end, snaps, col):
    """各站滾動 1 / 3 / 6 / 12 / 24 小時累積雨量的最大值;只算整段累積時間都在期間內的(起點 ≥ start),
    時間 = 快照時間(累積的結束時間)。"""
    n = len(col)
    out = {k: {"val": np.full(n, np.nan), "at": [None] * n} for k in RAIN_MAX}
    for t, f, rows in snaps:
        for k in RAIN_MAX:
            if k not in f or to_stamp(stamp_dt(t) - timedelta(minutes=RAIN_WIN[k])) < start:
                continue
            o = out[k]
            for r in rows:
                x = r[f[k]]
                if x is None:
                    continue
                i = col[r[f["id"]]]
                if np.isnan(o["val"][i]) or x > o["val"][i]:
                    o["val"][i], o["at"][i] = x, t
    return out


# ── 畫圖 ──
def load_shp(shp_dir):
    """底圖:使用者提供的 shp/TWN_CITY(縣市多邊形,WGS84 經緯度,含金馬、澎湖、離島)。
    陸地填色、縣市界、雨量內插的陸地範圍都用它。"""
    import shapefile
    r = shapefile.Reader(str(shp_dir / "TWN_CITY"), encoding="utf-8")
    geoms = [shapely.make_valid(shape(sh.__geo_interface__)).simplify(0.0003) for sh in r.shapes()]
    return {"city": geoms}


def land_mask(geo):
    """IDW 格點上的陸地範圍(依 TWN_CITY)。"""
    lon2, lat2 = np.meshgrid(idw.LON, idw.LAT)
    mask = np.zeros(lon2.shape, dtype=bool)
    for g in geo["city"]:
        mask |= shapely.contains_xy(g, lon2, lat2)
    return mask


def _polys(g):
    for part in getattr(g, "geoms", [g]):
        if part.geom_type == "Polygon":
            yield part
        elif hasattr(part, "geoms"):
            yield from _polys(part)


def draw_base(ax, geo, extent):
    """底圖:白色海洋、淺灰陸地(與網頁淺色主題相同);陸地只用台灣縣市界(NCDR)。"""
    x0, x1, y0, y1 = extent
    ax.set_xlim(x0, x1)
    ax.set_ylim(y0, y1)
    ax.set_aspect(1 / np.cos(np.deg2rad((y0 + y1) / 2)))
    ax.set_facecolor(LIGHT["bg"])
    ax.set_xticks([])
    ax.set_yticks([])
    patches = [MplPolygon(np.asarray(p.exterior.coords), closed=True) for g in geo["city"] for p in _polys(g)]
    ax.add_collection(PatchCollection(patches, facecolor=LIGHT["land"], edgecolor="none", zorder=0))


def draw_overlay(ax, geo):
    """縣市界(含海岸)疊在資料上面。"""
    for g in geo["city"]:
        for p in _polys(g):
            for ring in [p.exterior, *p.interiors]:
                c = np.asarray(ring.coords)
                ax.plot(c[:, 0], c[:, 1], color=LIGHT["county"], lw=0.5, zorder=5)


def beaufort(ms):
    return int(np.searchsorted(BEAUFORT, ms, side="right") - 1)


def make_figure(geo):
    """純地圖:沒有標題、文字、外框;金門、馬祖小圖放地圖左上海面;色階在右邊。"""
    fig = plt.figure(figsize=(6.6, 8.4), dpi=200)
    fig.patch.set_facecolor("white")
    ax = fig.add_axes([0.01, 0.01, 0.80, 0.98])
    draw_base(ax, geo, MAIN)
    ax.axis("off")
    insets = {}
    for name, rect in zip(INSETS, ([0.07, 0.77, 0.30, 0.20], [0.07, 0.585, 0.30, 0.17])):
        iax = ax.inset_axes(rect)
        draw_base(iax, geo, INSETS[name])
        iax.set_zorder(10)
        for sp in iax.spines.values():
            sp.set_edgecolor("#888")
        insets[name] = iax
    return fig, ax, insets


def finish(ax, insets, geo):
    for a in [ax, *insets.values()]:
        draw_overlay(a, geo)


def colorbar(fig, levels, colors, unit, cell_text=None):
    """右側直式色階:一格一色、等間距,刻度 = 各級下界,上方只標單位;cell_text = 格子裡的字(風級)。
    顏色比級數多一個 → 最後一色是「最後一個刻度以上」,畫成上方同樣大小的方格(例如 61.3 以上的 >17)。"""
    extend = len(levels) == len(colors)
    n = len(levels) - 1                                    # 刻度之間的格子數
    H = 0.74
    h = H * n / (n + 1) if extend else H                   # 加了上方格子,總高度不變
    cax = fig.add_axes([0.855, 0.04, 0.035, h])
    cmap = ListedColormap(colors)
    norm = BoundaryNorm(levels, cmap.N, extend="max" if extend else "neither")
    cb = fig.colorbar(ScalarMappable(norm=norm, cmap=cmap), cax=cax, ticks=levels, spacing="uniform",
                      extendrect=True, extendfrac=1 / n if extend else None)
    cb.ax.tick_params(labelsize=10, length=0)
    cb.ax.set_yticklabels([f"{x:g}" for x in levels])
    cb.outline.set_linewidth(0.5)
    cell_pt = (H - h) * fig.get_figheight() * 72          # 上方格子高度(點)
    cax.set_title(unit, fontsize=12, pad=8 + cell_pt)
    if cell_text:
        for i, t in enumerate(cell_text):
            cax.text(0.5, (i + 0.5) / n, t, transform=cax.transAxes, ha="center", va="center", fontsize=8.5,
                     fontweight="bold", color=text_color(colors[i]), clip_on=False)


def text_color(hex_):
    r, g, b = (int(hex_[i:i + 2], 16) for i in (1, 3, 5))
    return "#fff" if 0.299 * r + 0.587 * g + 0.114 * b < 140 else "#222"


def save(fig, path):
    """存檔;檔案被看圖軟體開著(Windows 會鎖住)時重試,還是不行就改存成 _new 檔名。"""
    import time
    for _ in range(3):
        try:
            fig.savefig(path, facecolor="white")
            break
        except OSError:
            time.sleep(1)
    else:
        alt = path.with_name(path.stem + "_new" + path.suffix)
        fig.savefig(alt, facecolor="white")
        print(f"  {path.name} 被其他程式開著,改存 {alt.name}")
    plt.close(fig)


def plot_rain(path, geo, stations, total):
    fig, ax, insets = make_figure(geo)
    ok = np.isfinite(total)
    lon = np.array([s["lon"] for s in stations])[ok]
    lat = np.array([s["lat"] for s in stations])[ok]
    grid = idw.interpolate(lon, lat, total[ok], land_mask(geo))
    grid = np.where(grid >= RAIN_LEVELS[0], grid, np.nan)
    cmap = ListedColormap(RAIN_COLORS)
    norm = BoundaryNorm(RAIN_LEVELS, cmap.N)
    for a in [ax, *insets.values()]:
        a.contourf(idw.LON, idw.LAT, np.minimum(grid, RAIN_LEVELS[-1] - 0.01), levels=RAIN_LEVELS, cmap=cmap,
                   norm=norm, zorder=2)
        a.scatter(lon, lat, s=1.2, c="#333", lw=0, zorder=6)
    finish(ax, insets, geo)
    colorbar(fig, RAIN_LEVELS, RAIN_COLORS, "mm")
    save(fig, path)


def plot_dots(path, geo, stations, v, levels, colors, unit, cell_text=None, white_from=None, min_level=None):
    """測站圓點圖(同網頁測站圖層):依色階上色,級數越高點越大、畫在越上面;white_from 級以上加白邊;
    min_level:低於這個值畫小灰點(雨量)。"""
    fig, ax, insets = make_figure(geo)
    ok = np.isfinite(v)
    order = np.argsort(np.where(ok, v, -1))
    order = order[ok[order]]
    lv = np.array([max(int(np.searchsorted(levels, x, side="right") - 1), 0) if np.isfinite(x) else 0 for x in v])
    lv = np.minimum(lv, len(colors) - 1)
    scale = 17 / max(len(colors) - 1, 1)                  # 大小比照風級 0–17
    for a in [ax, *insets.values()]:
        for i in order:
            s = stations[i]
            if min_level is not None and v[i] < min_level:
                a.scatter(s["lon"], s["lat"], s=4, c="#9e9e9e", lw=0, alpha=0.7, zorder=6)
                continue
            k = lv[i]
            strong = white_from is not None and k >= white_from
            r = 1.4 + 0.25 * k * scale
            a.scatter(s["lon"], s["lat"], s=(2 * r) ** 2, c=colors[k], zorder=6 + k * 0.01,
                      edgecolors="#fff" if strong else "#333", linewidths=1.2 if strong else 0.35)
    finish(ax, insets, geo)
    colorbar(fig, levels, colors, unit, cell_text=cell_text)
    save(fig, path)


def plot_wind(path, geo, stations, w):
    plot_dots(path, geo, stations, w["val"], BEAUFORT, WIND_COLORS + [WIND_OVER], "m/s",
              cell_text=[str(k) for k in range(len(WIND_COLORS))] + [">17"], white_from=14)


def write_csv(path, stations, cols):
    """各站表(依第一個數值欄由大到小)。"""
    with open(path, "w", newline="", encoding="utf-8-sig") as fp:
        wr = csv.writer(fp)
        wr.writerow(["站號", "站名", "縣市", "鄉鎮", "海拔(m)", "經度", "緯度", *cols.keys()])
        rows = [[s["id"], s["name"], s["county"], s["town"], fmt_alt(s["alt"]), s["lon"], s["lat"],
                 *[c[i] for c in cols.values()]] for i, s in enumerate(stations)]
        rows.sort(key=lambda r: -(r[7] if isinstance(r[7], (int, float)) and r[7] == r[7] else -1))
        wr.writerows(rows)


# 縣市表的地區:蘭嶼、綠島、基隆嶼、龜山島獨立,不併入原縣市;順序由北到南、離島最後
REGION_ORDER = ["基隆市", "臺北市", "新北市", "桃園市", "新竹市", "新竹縣", "苗栗縣", "臺中市", "彰化縣", "南投縣",
                "雲林縣", "嘉義市", "嘉義縣", "臺南市", "高雄市", "屏東縣", "宜蘭縣", "花蓮縣", "臺東縣",
                "澎湖縣", "金門縣", "連江縣", "基隆嶼", "龜山島", "綠島", "蘭嶼"]
WIND_ALT_MAX = 500                                      # 縣市風的統計只取海拔低於這個高度的站
# 縣市風的統計另外排除的山區鄉鎮(海拔雖低於 500 m,但屬山區,例如陽明山周邊)
WIND_EXCLUDE_TOWNS = {("臺北市", "士林區"), ("臺北市", "北投區")}


def fmt_alt(a):
    return "" if a is None else (int(a) if float(a).is_integer() else a)


def region(s):
    name, town = s["name"] or "", s["town"] or ""
    if town == "蘭嶼鄉":
        return "蘭嶼"
    if town == "綠島鄉":
        return "綠島"
    if name.startswith("龜山島"):
        return "龜山島"
    if name.startswith("基隆嶼"):
        return "基隆嶼"
    return s["county"] or "其他"


def region_table(stations, vals, label, unit, at=None, alt_max=None, wind=False, exclude=()):
    """各地區(縣市 + 獨立離島)的最大值與測站;alt_max:只取海拔低於此值的站(海拔不明的不算)。
    回傳 (表頭, 各列, 各列的最大值或 None)。"""
    best = {}
    for i, s in enumerate(stations):
        v = vals[i]
        if not np.isfinite(v):
            continue
        if alt_max is not None and (s["alt"] is None or s["alt"] >= alt_max):
            continue
        if (s["county"], s["town"]) in exclude:
            continue
        r = region(s)
        if r not in best or v > vals[best[r]]:
            best[r] = i
    order = REGION_ORDER + sorted(set(best) - set(REGION_ORDER))
    head = ["地區", "鄉鎮", "站名", f"{label}({unit})"] + (["風級"] if wind else []) + (["發生時間"] if at else [])         + ["站號", "海拔(m)"]
    rows, values = [], []
    for r in order:
        if r not in best:
            if r in REGION_ORDER:                          # 沒有站的地區留空列
                rows.append([r] + [""] * (len(head) - 1))
                values.append(None)
            continue
        i = best[r]
        s, v = stations[i], vals[i]
        rows.append([r, s["town"], s["name"], round(float(v), 1)] + ([bf_label(v)] if wind else [])
                    + ([tw(at[i], "%m/%d %H:%M") if at[i] else ""] if at else [])
                    + [s["id"], fmt_alt(s["alt"])])
        values.append(float(v))
    return head, rows, values


def write_table_csv(path, head, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as fp:
        wr = csv.writer(fp)
        wr.writerow(head)
        wr.writerows(rows)


def write_xlsx(path, sheets):
    """地區表 Excel(每個表一個工作表);風的最大值、風級兩欄依蒲福風級色階上色,雨量最大值依雨量色階。
    sheets:[(工作表名, 表頭, 各列, 各列數值, 色階 'wind' | 'rain')]"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for name, head, rows, values, kind in sheets:
        ws = wb.create_sheet(name)
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        levels, colors = (BEAUFORT, WIND_COLORS + [WIND_OVER]) if kind == "wind" else (RAIN_LEVELS, RAIN_COLORS)
        cols = [3] + ([4] if kind == "wind" else [])      # 0 起算:最大值、風級
        for k, v in enumerate(values):
            if v is None:
                continue
            j = min(max(int(np.searchsorted(levels, v, side="right") - 1), 0), len(colors) - 1)
            hexc = colors[j]
            for c in cols:
                cell = ws.cell(row=k + 2, column=c + 1)
                cell.fill = PatternFill("solid", fgColor=hexc[1:].upper())
                cell.font = Font(bold=True, color="FFFFFF" if text_color(hexc) == "#fff" else "222222")
                cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:                             # 欄寬依內容(中文算兩格)
            w = max(sum(2 if ord(ch) > 255 else 1 for ch in str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = w + 2
        ws.freeze_panes = "A2"
    try:
        wb.save(path)
    except OSError:
        alt = path.with_name(path.stem + "_new" + path.suffix)
        wb.save(alt)
        print(f"  {path.name} 被其他程式開著,改存 {alt.name}")


def bf_label(ms):
    k = beaufort(ms)
    return ">17" if k > 17 else str(k)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--start", required=True, help='台灣時間,例如 "2026-10-22 00:00"')
    ap.add_argument("--end", help="台灣時間;省略 = 最新一筆快照")
    ap.add_argument("--shp", default=str(config.ROOT / "shp"), help="底圖 shapefile 資料夾(含 TWN_CITY)")
    ap.add_argument("--local", action="store_true", help="只用本機 data/stations,不從 R2 補(排程那台電腦可用)")
    a = ap.parse_args()
    start = to_stamp(parse_tw(a.start))
    end = to_stamp(parse_tw(a.end)) if a.end else None
    run(start, end, Path(a.shp), not a.local)


def set_font():
    from matplotlib import font_manager
    have = {f.name for f in font_manager.fontManager.ttflist}
    plt.rcParams["font.family"] = [f for f in FONT if f in have] or ["sans-serif"]


def resolve_end(end, use_r2):
    """end(UTC stamp 或 None = 最新)→ 不晚於它的最後一筆測站快照。"""
    times = snapshot_times(use_r2)
    end = end or times[-1]
    return max([t for t in times if t <= end], default=end)


def run(start, end=None, shp=config.ROOT / "shp", use_r2=True, out=None):
    """start / end:UTC stamp(end None = 最新快照);out 省略 = charts/{起}_{迄}/。
    回傳 dict(out, start, end, stations, total, wind, rmax, geo)給 event_report 寫摘要。"""
    set_font()
    end = resolve_end(end, use_r2)
    print(f"期間 {tw(start)} – {tw(end)}(台灣)")
    times, snaps, meta, col = collect(start, end, use_r2)
    print(f"快照 {len(snaps)} 筆(應有 {int((stamp_dt(end) - stamp_dt(start)).total_seconds() // 600)} 筆),測站 {len(col)}")
    stations = [None] * len(col)
    for sid, i in col.items():
        stations[i] = meta[sid]

    total, seg, _ = total_rain(start, end, snaps, col)
    wind = max_wind(start, end, snaps, col)
    rmax = max_rain(start, end, snaps, col)
    print(f"雨量:{int(np.isfinite(total).sum())} 站接得起來;最多 {np.nanmax(total):.1f} mm")

    geo = load_shp(Path(shp))
    out = Path(out) if out else OUT_DIR / f"{tw(start, '%Y%m%d%H%M')}_{tw(end, '%Y%m%d%H%M')}"
    out.mkdir(parents=True, exist_ok=True)
    plot_rain(out / "rain_total.png", geo, stations, total)
    plot_dots(out / "rain_total_station.png", geo, stations, total, RAIN_LEVELS, RAIN_COLORS, "mm",
              min_level=RAIN_LEVELS[0])
    plot_wind(out / "wind_max.png", geo, stations, wind["ws"])
    plot_wind(out / "gust_max.png", geo, stations, wind["gust"])
    fmt_at = lambda ts: [tw(t, "%m/%d %H:%M") if t else "" for t in ts]
    write_csv(out / "rain_total.csv", stations, {"總累積雨量(mm)": [round(x, 1) if np.isfinite(x) else "" for x in total],
                                                  "使用段數": [int(x) if np.isfinite(x) else "" for x in seg]})
    write_csv(out / "wind_max.csv", stations, {
        "最大平均風(m/s)": [x if np.isfinite(x) else "" for x in wind["ws"]["val"]],
        "風向(°)": [x if np.isfinite(x) else "" for x in wind["ws"]["dir"]], "發生時間": fmt_at(wind["ws"]["at"]),
        "最大陣風(m/s)": [x if np.isfinite(x) else "" for x in wind["gust"]["val"]],
        "陣風風向(°)": [x if np.isfinite(x) else "" for x in wind["gust"]["dir"]],
        "陣風時間": fmt_at(wind["gust"]["at"])})
    rcols = {}
    for k, label in RAIN_MAX.items():
        rcols[f"{label}(mm)"] = [x if np.isfinite(x) else "" for x in rmax[k]["val"]]
        rcols[f"{label}時間"] = fmt_at(rmax[k]["at"])
    write_csv(out / "rain_max.csv", stations, rcols)
    tables = [
        ("總累積雨量", "county_rain_total", region_table(stations, total, "總累積雨量", "mm"), "rain"),
        ("最大平均風", "county_wind_max", region_table(stations, wind["ws"]["val"], "最大平均風", "m/s",
                                                   at=wind["ws"]["at"], alt_max=WIND_ALT_MAX, wind=True, exclude=WIND_EXCLUDE_TOWNS), "wind"),
        ("最大陣風", "county_gust_max", region_table(stations, wind["gust"]["val"], "最大陣風", "m/s",
                                                 at=wind["gust"]["at"], alt_max=WIND_ALT_MAX, wind=True, exclude=WIND_EXCLUDE_TOWNS), "wind"),
    ] + [(label, f"county_{k}_max", region_table(stations, rmax[k]["val"], label, "mm", at=rmax[k]["at"]), "rain")
         for k, label in RAIN_MAX.items()]
    for _, fname, (head, rows, _), _ in tables:
        write_table_csv(out / f"{fname}.csv", head, rows)
    write_xlsx(out / "county_stats.xlsx", [(name, *t, kind) for name, _, t, kind in tables])
    print(f"輸出:{out}")
    return {"out": out, "start": start, "end": end, "stations": stations, "total": total, "wind": wind,
            "rmax": rmax, "geo": geo}


if __name__ == "__main__":
    main()
